"""Make many: one video per spreadsheet row, from a CSV, TSV or .xlsx list. No Manim import.

The person brings a file with a first row of column names, matches those columns to the
design's fields, and gets one export per row. Everything here is plain Python: reading
the file, suggesting matches, turning cells into the values a template validates, and
keeping the small batch history that the Exports page groups by.

Cells are text and nothing else. A spreadsheet formula is never evaluated (``data_only``
reads only a value Excel itself stored); dates are written the way a person says them
("14 June 2026"), and control characters are dropped.
"""

from __future__ import annotations

import base64
import binascii
import copy
import csv
import datetime as _datetime
import io
import json
import os
import re
import secrets
import threading
import unicodedata
from typing import Optional

from .errors import StudioError
from .storage import slugify, utc_now

MAX_ROWS = 200
MAX_COLUMNS = 40
MAX_BYTES = 2 * 1024 * 1024
MAX_SHEET_ROWS = 20000  # a sheet of nothing but blank rows must not hang the app
MAX_TITLE = 120
BATCH_ID = re.compile(r"^[0-9a-f]{12}$")
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_SPACES = re.compile(r"\s+")
_DELIMITERS = ",\t;"
_MONTHS = ("January", "February", "March", "April", "May", "June", "July", "August",
           "September", "October", "November", "December")
_RESERVED = {"con", "prn", "aux", "nul", *(f"com{i}" for i in range(1, 10)),
             *(f"lpt{i}" for i in range(1, 10))}
_TEXT_WIDGETS = {"text", "textarea"}
_NUMBER_WIDGETS = {"number", "integer", "duration", "sparkle"}


class BatchError(StudioError):
    """A list that cannot be read or used yet, with a plain message."""

    status = 422


def megabytes(size: int) -> str:
    value = size / (1024 * 1024)
    return f"{value:.1f} MB" if value < 10 else f"{value:.0f} MB"


# -- reading the list ------------------------------------------------------------------

def clean_cell(value) -> str:
    """Plain text: control characters gone, outer whitespace trimmed."""
    if value is None:
        return ""
    return _CONTROL.sub("", str(value)).strip()


def cell_text(value) -> str:
    """A spreadsheet cell as the person would write it (dates included)."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, _datetime.datetime):
        text = f"{value.day} {_MONTHS[value.month - 1]} {value.year}"
        if value.hour or value.minute:
            text += f", {value.hour:02d}:{value.minute:02d}"
        return text
    if isinstance(value, _datetime.date):
        return f"{value.day} {_MONTHS[value.month - 1]} {value.year}"
    if isinstance(value, _datetime.time):
        return f"{value.hour:02d}:{value.minute:02d}"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(value, ".10g")
    return clean_cell(value)


def _trim(cells: list) -> list:
    """Without trailing empty cells: sheets often report formatted empty columns."""
    end = len(cells)
    while end and not cells[end - 1]:
        end -= 1
    return cells[:end]


def _table_from_rows(raw_rows) -> tuple:
    table = []
    for raw in raw_rows:
        cells = _trim([clean_cell(value) for value in raw])
        if any(cells):
            table.append(cells)
        if len(table) > MAX_ROWS + 1:
            break
    return _check_table(table)


def _check_table(table: list) -> tuple:
    if not table:
        raise BatchError("That list is empty. Add a first row with column names, then one row per video.")
    width = max(len(row) for row in table)
    if width > MAX_COLUMNS:
        raise BatchError(f"That list has {width} columns. The most you can use at once is {MAX_COLUMNS}. "
                         "Leave out the columns you don't need.")
    rows = len(table) - 1
    if rows > MAX_ROWS:
        raise BatchError(f"That list has {rows} rows. The most you can make at once is {MAX_ROWS}. "
                         "Split it into smaller lists.")
    if rows == 0:
        raise BatchError("That list has column names but no rows below them.")
    headers = _headers(table[0], width)
    body = [row + [""] * (width - len(row)) for row in table[1:]]
    return headers, body


def _headers(first: list, width: int) -> list:
    headers, seen = [], {}
    for index in range(width):
        name = clean_cell(first[index] if index < len(first) else "")
        name = _SPACES.sub(" ", name) or f"Column {index + 1}"
        count = seen.get(name.casefold(), 0) + 1
        seen[name.casefold()] = count
        headers.append(name if count == 1 else f"{name} ({count})")
    return headers


def decode_text(data: bytes) -> str:
    """UTF-8 with or without a BOM, else cp1252 (the encoding Excel writes on Windows)."""
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise BatchError("This file isn't readable as text. Save it as CSV or send the spreadsheet as .xlsx.")


def _delimiter(text: str, filename: str) -> str:
    if filename.lower().endswith(".tsv"):
        return "\t"
    # Text pasted from Excel is tab-separated, and its cells often contain commas
    # ("Smith, John<TAB>Paid"). A tab on the first real line is the stronger signal.
    first = next((line for line in text.splitlines() if line.strip()), "")
    if "\t" in first:
        return "\t"
    sample = "\n".join(text.splitlines()[:8])
    try:
        return csv.Sniffer().sniff(sample, delimiters=_DELIMITERS).delimiter
    except csv.Error:
        counts = {delimiter: first.count(delimiter) for delimiter in _DELIMITERS}
        return max(_DELIMITERS, key=lambda delimiter: counts[delimiter]) if any(counts.values()) else ","


def parse_text(text: str, filename: str = "") -> tuple:
    """(headers, rows of strings) from pasted or decoded text. The first row names the columns."""
    text = text.lstrip("\ufeff").replace("\r\n", "\n").replace("\r", "\n")
    if not text.strip():
        raise BatchError("Paste your list or choose a file first.")
    delimiter = _delimiter(text, filename)
    try:
        raw = list(csv.reader(io.StringIO(text), delimiter=delimiter, skipinitialspace=True))
    except csv.Error:
        raise BatchError("This list can't be read. Check that every quoted cell is closed.") from None
    return _table_from_rows(raw)


def _xlsx_table(data: bytes) -> tuple:
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover - the desktop runtime installs it
        raise BatchError("This app can't read Excel files yet. Save the sheet as CSV and try again.") from None
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - anything openpyxl refuses is a file we cannot read
        raise BatchError("This spreadsheet can't be read. Open it and save it again as .xlsx, or use CSV.") from None
    raw = []
    try:
        try:
            sheet = workbook.worksheets[0] if workbook.worksheets else None
            if sheet is None:
                raise BatchError("This spreadsheet has no sheets.")
            non_blank = 0
            for number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
                cells = [cell_text(value) for value in values]
                raw.append(cells)
                if any(cells):
                    non_blank += 1
                    if non_blank > MAX_ROWS + 1:  # enough kept rows to explain "too many rows"
                        break
                if number >= MAX_SHEET_ROWS:
                    break
        except BatchError:
            raise
        except Exception:  # noqa: BLE001 - a sheet can also break halfway through reading
            raise BatchError("This spreadsheet can't be read. Open it and save it again as .xlsx, or use CSV.") from None
    finally:
        workbook.close()
    return _table_from_rows(raw)


def parse_bytes(data: bytes, filename: str = "") -> tuple:
    """(headers, rows) from an uploaded .xlsx, CSV or TSV file."""
    if not data:
        raise BatchError("That file is empty. Choose a CSV, TSV or .xlsx file with your list.")
    if len(data) > MAX_BYTES:
        raise BatchError(f"That file is {megabytes(len(data))}. The most you can use at once is "
                         f"{MAX_BYTES // (1024 * 1024)} MB.")
    lower = filename.lower()
    if lower.endswith((".xls", ".xlsm", ".ods")) and not lower.endswith(".xlsx"):
        raise BatchError("Only .xlsx spreadsheets can be read. Open this file and save it as .xlsx or CSV.")
    if lower.endswith(".xlsx") or data.startswith(b"PK\x03\x04"):
        return _xlsx_table(data)
    return parse_text(decode_text(data), filename)


def decode_upload(value) -> bytes:
    """The bytes of a base64 (or data-URL) upload."""
    if not isinstance(value, str) or not value.strip():
        raise BatchError("Choose a CSV, TSV or .xlsx file.")
    text = value.strip()
    if text.startswith("data:") and "," in text:
        text = text.split(",", 1)[1]
    try:
        return base64.b64decode(text, validate=True)
    except (ValueError, binascii.Error):
        raise BatchError("That file couldn't be read. Choose it again.") from None


# -- which fields a list can fill --------------------------------------------------------

def _resolve(schema: dict, value: dict) -> dict:
    definitions = schema.get("$defs") or {}
    for _ in range(4):
        if not isinstance(value, dict) or "$ref" not in value:
            break
        value = definitions.get(str(value["$ref"]).rsplit("/", 1)[-1]) or {}
    return value if isinstance(value, dict) else {}


def mappable_fields(schema: dict) -> list:
    """The simple top-level text, number and color fields a column can fill.

    Lists, tables, pictures, icons, toggles and fixed choices are left out: a column of
    text cannot honestly become one of those, and a wrong guess would only be a problem
    the person has to fix later.
    """
    properties = (schema or {}).get("properties") or {}
    required = set((schema or {}).get("required") or [])
    fields = []
    for name, raw in properties.items():
        prop = _resolve(schema, raw or {})
        variants = prop.get("anyOf") or prop.get("oneOf") or []
        string_variant = next((variant for variant in variants
                               if isinstance(variant, dict) and variant.get("type") == "string"), None)
        if string_variant:
            prop = {**prop, **string_variant}
        widget = prop.get("x-widget")
        type_ = prop.get("type")
        enum = prop.get("enum")
        if widget == "color":
            kind = "color"
        elif widget in _TEXT_WIDGETS:
            kind = "text"
        elif widget in _NUMBER_WIDGETS and not prop.get("x-choice-of"):
            kind = "number"
        elif widget in ("image", "icon", "list", "table", "toggle", "choice") or prop.get("x-choice-of"):
            continue
        elif type_ == "string" and not enum:
            kind = "text"
        elif type_ in ("number", "integer") and not enum:
            kind = "number"
        else:
            continue
        fields.append({
            "field": name,
            "label": prop.get("title") or name.replace("_", " ").capitalize(),
            "kind": kind,
            "max_length": prop.get("maxLength") if kind == "text" else None,
            "required": name in required,
        })
    return fields


def normalize_key(text: str) -> str:
    """For matching only: case, spaces, punctuation and accents don't matter."""
    ascii_text = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]", "", ascii_text.casefold())


def suggest_mapping(fields: list, headers: list) -> dict:
    """field name -> column index, one column per field, by field name and label."""
    normalized = [normalize_key(header) for header in headers]
    pairs = []
    for order, field in enumerate(fields):
        names = [normalize_key(field["field"]), normalize_key(field["label"])]
        for index, column in enumerate(normalized):
            if not column:
                continue
            best = 0
            for name in names:
                if not name:
                    continue
                if column == name:
                    best = max(best, 100)
                elif column.startswith(name) or name.startswith(column):
                    best = max(best, 70)
                elif name in column or column in name:
                    shorter, longer = sorted((len(name), len(column)))
                    ratio = shorter / max(1, longer)
                    if ratio >= 0.6:
                        best = max(best, 60)
            if best:
                pairs.append((best, order, index))
    pairs.sort(key=lambda pair: (-pair[0], pair[1], pair[2]))
    mapping, used = {}, set()
    for score, order, index in pairs:
        field, column = fields[order]["field"], headers[index]
        if field in mapping or column in used:
            continue
        mapping[field] = index
        used.add(column)
    return mapping


def suggest_title_column(fields: list, headers: list, mapping: dict) -> Optional[int]:
    """A likely column for naming each video: "Title", "Name", a headline, else the first."""
    by_field = {field["field"]: field for field in fields}
    preferred = []
    for field_name, index in mapping.items():
        field = by_field.get(field_name) or {}
        name = normalize_key(field_name)
        label = normalize_key(field.get("label", ""))
        if name in {"title", "name", "videotitle", "filename"} or label in {"title", "name"}:
            rank = 0
        elif "title" in name or "title" in label:
            rank = 1
        elif "headline" in name or "headline" in label:
            rank = 2
        elif "name" in name or "name" in label:
            rank = 3
        else:
            continue
        preferred.append((rank, index))
    if preferred:
        return min(preferred)[1]
    for index, header in enumerate(headers):
        if normalize_key(header) in {"title", "name", "filename", "videotitle", "file"}:
            return index
    return 0 if headers else None


# -- one row's values ---------------------------------------------------------------------

def normalize_color(text: str) -> str:
    value = text.strip()
    if re.fullmatch(r"[0-9a-fA-F]{6}", value):
        return "#" + value.upper()
    if re.fullmatch(r"#[0-9a-fA-F]{6}", value):
        return value.upper()
    return value


def _cell(row: list, index) -> str:
    if isinstance(index, bool) or not isinstance(index, int) or index < 0 or index >= len(row):
        return ""
    return clean_cell(row[index])


def row_params(base_params: dict, fields: list, mapping: dict, row: list) -> dict:
    """The row's values on top of the design's own text; an empty cell keeps the design."""
    params = copy.deepcopy(base_params or {})
    by_field = {field["field"]: field for field in fields}
    for name, index in mapping.items():
        field = by_field.get(name)
        if field is None:
            continue
        value = _cell(row, index)
        if not value:
            continue
        params[name] = normalize_color(value) if field["kind"] == "color" else value
    return params


def row_title(row: list, column, fallback: str = "") -> str:
    """The title column's text, made safe to show and use as a file-name stem."""
    text = _cell(row, column) if column is not None else ""
    text = _SPACES.sub(" ", text).strip()
    return text[:MAX_TITLE] or fallback


def safe_file_stem(text: str, fallback: str = "video") -> str:
    """A readable file-name stem that cannot escape its folder or clash with Windows."""
    stem = slugify(text, fallback=fallback)
    first = stem.split(".", 1)[0].casefold()
    if first in _RESERVED:
        stem = f"{stem}-video"
    return stem or "video"


def plain_message(message: str) -> str:
    """Pydantic's sentence, said the way a person would say it."""
    text = str(message or "").strip()
    text = re.sub(r"^Value error,\s*", "", text)
    match = re.search(r"at most (\d+) characters", text)
    if match:
        return f"Use at most {match.group(1)} characters."
    match = re.search(r"at least (\d+) characters", text)
    if match:
        return f"Use at least {match.group(1)} characters."
    match = re.search(r"greater than or equal to (-?[\d.]+)", text)
    if match:
        return f"Use {match.group(1)} or more."
    match = re.search(r"less than or equal to (-?[\d.]+)", text)
    if match:
        return f"Use {match.group(1)} or less."
    if "Field required" in text or "missing" in text.casefold():
        return "This is needed."
    if re.search(r"valid number|valid integer|parse string as an integer", text, re.I):
        return "Use a number."
    if re.search(r"valid string", text, re.I):
        return "Use text."
    if re.search(r"valid boolean", text, re.I):
        return "Use yes or no."
    if re.search(r"valid color|like #2563EB", text, re.I):
        return "Use a color like #2563EB."
    if text.startswith("Input should be "):
        return "Use a " + text[len("Input should be "):] + "."
    return text or "This value can't be used."


def plain_errors(errors: list, fields: list) -> list:
    """Per-field problems in the words the person sees, keyed by the field's label."""
    by_field = {field["field"]: field for field in fields}
    problems = []
    for error in errors or []:
        path = str(error.get("field") or "")
        name = path.rsplit(".", 1)[-1]
        field = by_field.get(name) or {}
        label = field.get("label") or name.replace("_", " ").capitalize() or "This value"
        problems.append({"field": path, "label": label, "message": plain_message(error.get("message", ""))})
    return problems


# -- the batch history ---------------------------------------------------------------------

class BatchStore:
    """Small JSON history of batches, kept with the workspace like the export index."""

    MAX_RECORDS = 200
    MAX_FILE = 1024 * 1024

    def __init__(self, workspace: str):
        self.path = os.path.join(workspace, "studio", "batches.json")
        self._lock = threading.RLock()

    def _read_locked(self) -> list:
        try:
            with open(self.path, encoding="utf-8") as source:
                raw = source.read(self.MAX_FILE + 1)
            if len(raw) > self.MAX_FILE:
                return []
            data = json.loads(raw)
        except (OSError, ValueError):
            return []
        records = data.get("batches") if isinstance(data, dict) else None
        return [record for record in records if isinstance(record, dict)
                and BATCH_ID.match(str(record.get("id", "")))] if isinstance(records, list) else []

    def list(self) -> list:
        with self._lock:
            return self._read_locked()

    def get(self, batch_id: str) -> Optional[dict]:
        if not isinstance(batch_id, str) or not BATCH_ID.match(batch_id):
            return None
        for record in self.list():
            if record.get("id") == batch_id:
                return record
        return None

    def add(self, record: dict) -> dict:
        import file_ops

        with self._lock:
            records = [record, *self.list()][:self.MAX_RECORDS]
            data = json.dumps({"version": 1, "batches": records}, indent=1).encode("utf-8")
            os.makedirs(os.path.dirname(self.path), exist_ok=True)
            file_ops.atomic_write(self.path, data)
            return record

    def new_id(self) -> str:
        known = {record.get("id") for record in self.list()}
        for _ in range(100):
            batch_id = secrets.token_hex(6)
            if batch_id not in known:
                return batch_id
        raise StudioError("Could not pick a batch id.")  # pragma: no cover


def folder_name(name: str, batch_id: str) -> str:
    """The one folder a batch's files go to: its name, made safe, plus the batch id."""
    return f"{safe_file_stem(name, fallback='batch')}-{batch_id}"


def new_record(batch_id: str, name: str, template: str, title: str, total: int, folder: str,
               format: str, quality: str, row_errors: list) -> dict:
    return {"id": batch_id, "name": name, "template": template, "title": title, "created": utc_now(),
            "total": int(total), "folder": folder, "format": format, "quality": quality,
            "row_errors": row_errors}


STATES = ("queued", "rendering", "done", "failed", "cancelled")


def counts(records: list) -> dict:
    result = {state: 0 for state in STATES}
    for record in records:
        status = record.get("status")
        if status in result:
            result[status] += 1
    return result
